"""python test_ledger.py 로 실행한다. 먼저 make_sample.py로 samples/를 만들어 둔다."""
import csv
import os
import tempfile
from datetime import date

from openpyxl import load_workbook

import build_ledger as bl

HERE = os.path.dirname(os.path.abspath(__file__))
SAMPLE = os.path.join(HERE, "samples", "잔고_가상펀드.xlsx")
TRUTH = os.path.join(HERE, "samples", "정답_거래내역.csv")


def test_xirr_known_values():
    # 365일 뒤 10% 이익 → 10%
    r = bl.xirr([(date(2025, 1, 1), -1000), (date(2026, 1, 1), 1100)])
    assert abs(r - 0.10) < 1e-9, r
    # 엑셀 XIRR 도움말 예제 → 37.34%
    flows = [(date(2008, 1, 1), -10000), (date(2008, 3, 1), 2750), (date(2008, 10, 30), 4250),
             (date(2009, 2, 15), 3250), (date(2009, 4, 1), 2750)]
    assert abs(bl.xirr(flows) - 0.373362535) < 1e-6, bl.xirr(flows)
    assert bl.xirr([(date(2025, 1, 1), -1000)]) is None


def test_ledger_matches_truth():
    _, ledger, value_errors, _, skipped = bl.build(SAMPLE)
    got = sorted((d.isoformat(), inv, kind, int(units)) for d, inv, _, kind, units, _, _ in ledger)
    with open(TRUTH, encoding="utf-8-sig") as f:
        want = sorted((r["기준일"], r["수익자번호"], r["구분"], int(r["좌수변동"])) for r in csv.DictReader(f))
    assert got == want, set(got) ^ set(want)
    assert value_errors == [] and skipped == []


def test_detects_value_mismatch():
    wb = load_workbook(SAMPLE)
    wb.worksheets[3].cell(4, 5).value += 5000  # 한 수익자의 평가금액을 일부러 틀리게
    with tempfile.TemporaryDirectory() as tmp:
        bad = os.path.join(tmp, "bad.xlsx")
        wb.save(bad)
        _, _, value_errors, _, _ = bl.build(bad)
    assert len(value_errors) == 1, value_errors


def test_summary_balances():
    snaps, ledger, _, names, _ = bl.build(SAMPLE)
    summary, by_date, _ = bl.summarize(snaps, ledger, names)
    # 수익자별 (납입 - 환매) 합계 = 날짜별 순유입 합계
    assert abs(sum(r[2] - r[3] for r in summary) - sum(i - o for i, o in by_date.values())) < 1e-6
    # 현재 평가액 합계 = 마지막 시트 평가금액 합계 (수익자당 반올림 1원까지 허용)
    last_values = sum(v for _, _, _, v in snaps[-1][1].values())
    assert abs(sum(r[4] for r in summary) - last_values) <= len(summary)


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("OK", name)
