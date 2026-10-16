"""기준일별 잔고 시트만 있는 펀드 엑셀에서 수익자별 거래내역과 수익률(XIRR)을 만든다.

잔고 변화로 거래를 추정한다:
    처음 생김/0에서 증가 → 설정, 증가 → 추가설정, 감소 → 일부환매, 0이 됨 → 전액환매
거래금액 = 좌수변동 × 해당 기준일 기준가격 / 1000

한계: 두 기준일 사이에 여러 번 거래했으면 순변동 1건으로만 보인다.

사용법:
    python build_ledger.py samples/잔고_가상펀드.xlsx -o 거래원장.xlsx
"""
import argparse
import re
import sys
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

HEADERS = {"id": ("수익자번호", "계좌번호", "고객번호"), "name": ("수익자명", "고객명", "성명"), "units": ("보유좌수", "좌수", "잔고좌수"),
           "nav": ("기준가격", "기준가"), "value": ("평가금액", "평가액")}


def sheet_date(ws):
    m = re.fullmatch(r"(\d{4})[-.]?(\d{2})[-.]?(\d{2})", ws.title.strip())
    if m:
        return date(*map(int, m.groups()))
    for row in ws.iter_rows(max_row=5, values_only=True):
        for cell in row:
            m = re.search(r"기준일\s*[:：]?\s*(\d{4})[-.](\d{2})[-.](\d{2})", str(cell or ""))
            if m:
                return date(*map(int, m.groups()))
    return None


def read_snapshot(ws):
    """헤더 줄을 찾아 {수익자번호: (이름, 좌수, 기준가, 평가금액)}과 기준가를 돌려준다."""
    rows = list(ws.iter_rows(values_only=True))
    for h, row in enumerate(rows):
        cells = [str(c or "").replace(" ", "") for c in row]
        cols = {k: next((i for i, c in enumerate(cells) if c in names), None) for k, names in HEADERS.items()}
        if cols["id"] is not None and cols["units"] is not None:
            break
    else:
        raise ValueError(f"[{ws.title}] 수익자번호/좌수 헤더를 찾지 못했습니다.")
    data, nav = {}, None
    for row in rows[h + 1:]:
        inv = row[cols["id"]]
        if inv is None or not re.fullmatch(r"\d{6,}", str(inv).strip()):  # 합계·빈 줄 제외
            continue
        units = float(row[cols["units"]] or 0)
        nav = float(row[cols["nav"]]) if cols["nav"] is not None else nav
        value = float(row[cols["value"]]) if cols["value"] is not None and row[cols["value"]] is not None else None
        data[str(inv).strip()] = (row[cols["name"]] if cols["name"] is not None else "", units, nav, value)
    return data, nav


def xirr(flows):
    """[(날짜, 금액)]의 연환산 내부수익률. 투자는 음수, 회수·평가액은 양수. 엑셀 XIRR과 같은 365일 기준."""
    if not any(a < 0 for _, a in flows) or not any(a > 0 for _, a in flows):
        return None
    d0 = min(d for d, _ in flows)

    def npv(r):
        return sum(a / (1 + r) ** ((d - d0).days / 365) for d, a in flows)

    lo, hi = -0.9999, 100.0
    if npv(lo) * npv(hi) > 0:
        return None
    for _ in range(200):
        mid = (lo + hi) / 2
        if npv(lo) * npv(mid) <= 0:
            hi = mid
        else:
            lo = mid
    return (lo + hi) / 2


def build(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    snaps, skipped = [], []
    for ws in wb.worksheets:
        d = sheet_date(ws)
        if d is None:
            skipped.append(ws.title)
            continue
        data, nav = read_snapshot(ws)
        snaps.append((d, data, nav))
    snaps.sort(key=lambda s: s[0])

    ledger, value_errors, prev, names = [], [], {}, {}
    for d, data, nav in snaps:
        for inv in sorted(set(prev) | set(data)):
            before = prev.get(inv, (None, 0))[1]
            name, after, row_nav, value = data.get(inv, (names.get(inv, ""), 0, nav, None))
            names[inv] = name
            if value is not None and abs(after * row_nav / 1000 - value) > 1:
                value_errors.append((d, inv, name, after * row_nav / 1000, value))
            delta = after - before
            if delta == 0:
                continue
            kind = "설정" if before == 0 else "추가설정" if delta > 0 else "전액환매" if after == 0 else "일부환매"
            ledger.append((d, inv, name, kind, delta, nav, delta * nav / 1000))
        prev = data
    return snaps, ledger, value_errors, names, skipped


def summarize(snaps, ledger, names):
    last_date, last_data, _ = snaps[-1]
    summary, fund_flows = [], []
    for inv in sorted(names):
        rows = [r for r in ledger if r[1] == inv]
        paid = sum(r[6] for r in rows if r[6] > 0)
        received = -sum(r[6] for r in rows if r[6] < 0)
        _, units, nav, _ = last_data.get(inv, ("", 0, None, None))
        current = units * nav / 1000 if units else 0
        flows = [(r[0], -r[6]) for r in rows] + ([(last_date, current)] if current else [])
        fund_flows += flows
        summary.append((inv, names[inv], paid, received, current, received + current - paid, xirr(flows)))
    by_date = {}
    for d, _, _, _, _, _, amount in ledger:
        flow = by_date.setdefault(d, [0, 0])
        flow[0 if amount > 0 else 1] += abs(amount)
    return summary, by_date, xirr(fund_flows)


def write_excel(path, ledger, summary, by_date, fund_irr, value_errors, skipped, snaps):
    wb = Workbook()
    ws = wb.active
    ws.title = "거래내역"
    ws.append(["기준일", "수익자번호", "수익자명", "구분(추정)", "좌수변동", "기준가격", "거래금액(원)"])
    for r in ledger:
        ws.append(list(r))

    ws2 = wb.create_sheet("수익자별 요약")
    ws2.append(["수익자번호", "수익자명", "총 납입", "총 환매", "현재 평가액", "손익", "XIRR"])
    for r in summary:
        ws2.append(list(r))
    ws2.append(["합계", None] + [sum(r[i] for r in summary) for i in range(2, 6)] + [fund_irr])

    ws3 = wb.create_sheet("펀드 현금흐름")
    ws3.append(["기준일", "설정 유입", "환매 유출", "순유입"])
    for d in sorted(by_date):
        inflow, outflow = by_date[d]
        ws3.append([d, inflow, outflow, inflow - outflow])

    ws4 = wb.create_sheet("검증")
    ws4.append(["평가금액 불일치 (좌수×기준가/1000 vs 원본, 1원 초과 차이)"])
    ws4.append(["기준일", "수익자번호", "수익자명", "계산값", "원본값"])
    for r in value_errors:
        ws4.append(list(r))
    if not value_errors:
        ws4.append(["불일치 없음"])
    ws4.append([])
    ws4.append(["읽은 시트", f"{len(snaps)}개 ({snaps[0][0]} ~ {snaps[-1][0]})"])
    ws4.append(["날짜를 몰라 건너뛴 시트", ", ".join(skipped) or "없음"])

    fill = PatternFill("solid", fgColor="DDE7F3")
    for sheet in wb.worksheets:
        for c in sheet[1]:
            c.font = Font(bold=True)
            c.fill = fill
        for col in "ABCDEFG":
            sheet.column_dimensions[col].width = 16
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, (date, datetime)):
                    c.number_format = "yyyy-mm-dd"
                elif isinstance(c.value, float):
                    c.number_format = "0.00%" if sheet.cell(1, c.column).value == "XIRR" else "#,##0;[Red]-#,##0"
        sheet.freeze_panes = "A2"
    wb.save(path)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("-o", "--output", default="거래원장.xlsx")
    args = ap.parse_args(argv)

    snaps, ledger, value_errors, names, skipped = build(args.xlsx)
    summary, by_date, fund_irr = summarize(snaps, ledger, names)
    write_excel(args.output, ledger, summary, by_date, fund_irr, value_errors, skipped, snaps)
    print(f"시트 {len(snaps)}개, 수익자 {len(names)}명, 거래 {len(ledger)}건, 평가금액 불일치 {len(value_errors)}건")
    print(f"펀드 전체 XIRR {fund_irr:.2%} -> {args.output}" if fund_irr is not None else f"-> {args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
